"""
企业级文档摄取管道
Enterprise Document Ingestion Pipeline

基于业界最佳实践设计:
- LlamaIndex IngestionPipeline 架构
- 多种分块策略（语义分块、递归分块、层次化分块）
- 异步处理队列
- 进度跟踪与错误重试
- 文件去重与增量处理
- 批量并发处理
"""

import os
import re
import json
import uuid
import hashlib
import asyncio
import time
import tempfile
from typing import List, Dict, Optional, Any, Callable, Generator, AsyncGenerator
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from abc import ABC, abstractmethod
from concurrent.futures import ThreadPoolExecutor
import threading
from loguru import logger

from src.common.domain_profile_service import get_domain_profile_service
from src.common.schema_category_resolver import infer_category_from_text
from src.common.industry_schema_service import get_industry_schema_service
from src.infrastructure.runtime_paths import get_data_dir
from src.rag.knowledge_understanding_service import get_knowledge_understanding_service


class ProcessingStatus(Enum):
    """处理状态枚举"""
    PENDING = "pending"
    UPLOADING = "uploading"
    PARSING = "parsing"
    CHUNKING = "chunking"
    EMBEDDING = "embedding"
    STORING = "storing"
    OCR_REQUIRED = "ocr_required"
    COMPLETED = "completed"
    FAILED = "failed"
    CANCELLED = "cancelled"


class ChunkStrategy(Enum):
    """分块策略枚举"""
    FIXED_SIZE = "fixed_size"
    RECURSIVE = "recursive"
    SEMANTIC = "semantic"
    HIERARCHICAL = "hierarchical"
    SLIDING_WINDOW = "sliding_window"


QA_SINGLE_CHUNK_MAX_BYTES = 512


class OCRRequiredError(RuntimeError):
    """文档需要先经过 OCR 才能继续摄取。"""


@dataclass
class ProcessingProgress:
    """处理进度"""
    taskId: str
    fileName: str
    status: ProcessingStatus
    progress: float = 0.0
    currentStep: str = ""
    totalSteps: int = 5
    completedSteps: int = 0
    errorMessage: Optional[str] = None
    startTime: Optional[datetime] = None
    endTime: Optional[datetime] = None
    documentId: str = ""
    contentHash: str = ""
    chunkCount: int = 0
    vectorCount: int = 0
    chunkIds: List[str] = field(default_factory=list)
    knowledgeItemIds: List[str] = field(default_factory=list)
    indexLedger: Dict = field(default_factory=dict)
    chunkQualitySummary: Dict = field(default_factory=dict)
    
    def toDict(self) -> Dict:
        """转换为字典"""
        return {
            "taskId": self.taskId,
            "fileName": self.fileName,
            "status": self.status.value,
            "progress": self.progress,
            "currentStep": self.currentStep,
            "totalSteps": self.totalSteps,
            "completedSteps": self.completedSteps,
            "errorMessage": self.errorMessage,
            "startTime": self.startTime.isoformat() if self.startTime else None,
            "endTime": self.endTime.isoformat() if self.endTime else None,
            "documentId": self.documentId,
            "contentHash": self.contentHash,
            "chunkCount": self.chunkCount,
            "vectorCount": self.vectorCount,
            "chunkIds": self.chunkIds,
            "knowledgeItemIds": self.knowledgeItemIds,
            "indexLedger": self.indexLedger,
            "chunkQualitySummary": self.chunkQualitySummary,
            "duration": (self.endTime - self.startTime).total_seconds() if self.startTime and self.endTime else None
        }


@dataclass
class DocumentChunk:
    """文档分块"""
    id: str
    content: str
    documentId: str
    documentName: str
    chunkIndex: int
    chunkType: str = "text"
    parentChunkId: Optional[str] = None
    childChunkIds: List[str] = field(default_factory=list)
    metadata: Dict = field(default_factory=dict)
    embedding: Optional[List[float]] = None
    createdAt: Optional[datetime] = None
    
    def __post_init__(self):
        if self.createdAt is None:
            self.createdAt = datetime.now()


@dataclass
class ProcessedDocument:
    """处理后的文档"""
    id: str
    name: str
    contentHash: str
    fileType: str
    fileSize: int
    chunks: List[DocumentChunk]
    metadata: Dict = field(default_factory=dict)
    createdAt: Optional[datetime] = None
    
    def __post_init__(self):
        if self.createdAt is None:
            self.createdAt = datetime.now()


class BaseChunker(ABC):
    """分块器基类"""
    
    @abstractmethod
    def chunk(self, text: str, documentId: str, documentName: str, metadata: Optional[Dict] = None) -> List[DocumentChunk]:
        """分块方法"""
        pass


class FixedSizeChunker(BaseChunker):
    """
    固定大小分块器
    
    特点：简单高效，适合结构简单的文档
    优化：增大分块大小到800字符，重叠15%
    """
    
    def __init__(self, chunkSize: int = 800, overlap: int = 120):
        self.chunkSize = max(int(chunkSize or 800), 100)
        self.overlap = max(0, min(int(overlap or 0), self.chunkSize - 1))
    
    def chunk(self, text: str, documentId: str, documentName: str, metadata: Optional[Dict] = None) -> List[DocumentChunk]:
        """按固定大小分块"""
        metadata = metadata or {}
        chunks = []
        text = text or ""
        start = 0
        chunkIndex = 0

        step = max(self.chunkSize - self.overlap, 1)
        while start < len(text):
            end = min(start + self.chunkSize, len(text))
            chunkText = text[start:end].strip()

            if len(chunkText.strip()) >= 50:
                chunkId = self._generateChunkId(documentId, chunkIndex)
                chunks.append(DocumentChunk(
                    id=chunkId,
                    content=chunkText,
                    documentId=documentId,
                    documentName=documentName,
                    chunkIndex=chunkIndex,
                    chunkType="fixed_size",
                    metadata={
                        **metadata,
                        "charCount": len(chunkText),
                        "wordCount": len(chunkText.split()),
                        "startPos": start,
                        "endPos": end
                    }
                ))
                chunkIndex += 1

            if end >= len(text):
                break
            start += step
        
        return chunks
    
    def _generateChunkId(self, documentId: str, chunkIndex: int) -> str:
        """生成分块ID"""
        content = f"{documentId}_{chunkIndex}_{time.time()}"
        return hashlib.md5(content.encode()).hexdigest()[:16]


class RecursiveChunker(BaseChunker):
    """
    递归字符分块器
    
    特点：按分隔符优先级递归分割，保持语义完整性
    参考：LangChain RecursiveCharacterTextSplitter
    优化：增大分块大小到800字符，重叠15%
    """
    
    def __init__(self, chunkSize: int = 800, overlap: int = 120):
        self.chunkSize = chunkSize
        self.overlap = overlap
        
        self.separators = [
            "\n\n\n",
            "\n\n",
            "\n",
            "。",
            "！",
            "？",
            "；",
            "，",
            " ",
            ""
        ]
    
    def chunk(self, text: str, documentId: str, documentName: str, metadata: Optional[Dict] = None) -> List[DocumentChunk]:
        """递归分块"""
        metadata = metadata or {}
        chunks = []
        
        splits = self._splitText(text)
        
        currentChunk = []
        currentSize = 0
        chunkIndex = 0
        
        for split in splits:
            splitSize = len(split)
            
            if currentSize + splitSize > self.chunkSize and currentChunk:
                chunkText = "".join(currentChunk)
                if len(chunkText.strip()) >= 50:
                    chunkId = self._generateChunkId(documentId, chunkIndex)
                    chunks.append(DocumentChunk(
                        id=chunkId,
                        content=chunkText,
                        documentId=documentId,
                        documentName=documentName,
                        chunkIndex=chunkIndex,
                        chunkType="recursive",
                        metadata={
                            **metadata,
                            "charCount": len(chunkText),
                            "separator": self._detectSeparator(chunkText)
                        }
                    ))
                    chunkIndex += 1
                
                overlapText = self._getOverlap(currentChunk)
                currentChunk = [overlapText] if overlapText else []
                currentSize = len(overlapText) if overlapText else 0
            
            currentChunk.append(split)
            currentSize += splitSize
        
        if currentChunk:
            chunkText = "".join(currentChunk)
            if len(chunkText.strip()) >= 50:
                chunkId = self._generateChunkId(documentId, chunkIndex)
                chunks.append(DocumentChunk(
                    id=chunkId,
                    content=chunkText,
                    documentId=documentId,
                    documentName=documentName,
                    chunkIndex=chunkIndex,
                    chunkType="recursive",
                    metadata={
                        **metadata,
                        "charCount": len(chunkText)
                    }
                ))
        
        return chunks
    
    def _splitText(self, text: str) -> List[str]:
        """按分隔符递归分割"""
        if len(text) <= self.chunkSize:
            return [text] if text.strip() else []
        
        for separator in self.separators:
            if separator in text:
                splits = text.split(separator)
                result = []
                for split in splits:
                    if split.strip():
                        subSplits = self._splitText(split)
                        result.extend(subSplits)
                return result
        
        return [text]
    
    def _getOverlap(self, chunks: List[str]) -> str:
        """获取重叠文本"""
        if not chunks:
            return ""
        
        combined = "".join(chunks)
        if len(combined) <= self.overlap:
            return combined
        
        return combined[-self.overlap:]
    
    def _detectSeparator(self, text: str) -> str:
        """检测主要分隔符"""
        for sep in self.separators:
            if sep in text:
                return sep.replace("\n", "\\n")
        return "none"
    
    def _generateChunkId(self, documentId: str, chunkIndex: int) -> str:
        """生成分块ID"""
        content = f"{documentId}_{chunkIndex}_{time.time()}"
        return hashlib.md5(content.encode()).hexdigest()[:16]


class SemanticChunker(BaseChunker):
    """
    语义分块器
    
    特点：基于句子语义相似度进行分块，保持语义连贯性
    参考：LlamaIndex SemanticSplitter
    优化：降低相似度阈值到0.6，增大最小分块到200字符
    """
    
    def __init__(
        self,
        embeddingService=None,
        similarityThreshold: float = 0.6,
        minChunkSize: int = 200,
        targetChunkSize: int = 500
    ):
        self.embeddingService = embeddingService
        self.similarityThreshold = similarityThreshold
        self.minChunkSize = max(int(minChunkSize or 200), 50)
        self.targetChunkSize = max(int(targetChunkSize or 500), self.minChunkSize)
    
    def chunk(self, text: str, documentId: str, documentName: str, metadata: Optional[Dict] = None) -> List[DocumentChunk]:
        """语义分块"""
        metadata = metadata or {}
        
        sentences = self._splitSentences(text)
        
        if not sentences:
            return []
        
        if self.embeddingService:
            return self._semanticChunk(sentences, documentId, documentName, metadata)
        else:
            return self._fallbackChunk(sentences, documentId, documentName, metadata)
    
    def _splitSentences(self, text: str) -> List[str]:
        """分割句子"""
        pattern = r'(?<=[。！？\n])'
        sentences = re.split(pattern, text)
        return [s.strip() for s in sentences if s.strip()]
    
    def _semanticChunk(self, sentences: List[str], documentId: str, documentName: str, metadata: Dict) -> List[DocumentChunk]:
        """基于语义相似度分块"""
        chunks = []
        currentSentences = []
        chunkIndex = 0
        
        try:
            embeddings = self.embeddingService.embed(sentences)
        except Exception as e:
            logger.warning(f"语义嵌入失败，使用回退策略: {e}")
            return self._fallbackChunk(sentences, documentId, documentName, metadata)
        
        for i, sentence in enumerate(sentences):
            currentSentences.append(sentence)
            
            if i < len(sentences) - 1:
                similarity = self._cosineSimilarity(embeddings[i], embeddings[i + 1])
                
                if similarity < self.similarityThreshold:
                    chunkText = "".join(currentSentences)
                    if len(chunkText) >= self.minChunkSize:
                        chunkId = self._generateChunkId(documentId, chunkIndex)
                        chunks.append(DocumentChunk(
                            id=chunkId,
                            content=chunkText,
                            documentId=documentId,
                            documentName=documentName,
                            chunkIndex=chunkIndex,
                            chunkType="semantic",
                            metadata={
                                **metadata,
                                "charCount": len(chunkText),
                                "sentenceCount": len(currentSentences),
                                "avgSimilarity": similarity
                            }
                        ))
                        chunkIndex += 1
                        currentSentences = []
        
        if currentSentences:
            chunkText = "".join(currentSentences)
            if len(chunkText) >= self.minChunkSize:
                chunkId = self._generateChunkId(documentId, chunkIndex)
                chunks.append(DocumentChunk(
                    id=chunkId,
                    content=chunkText,
                    documentId=documentId,
                    documentName=documentName,
                    chunkIndex=chunkIndex,
                    chunkType="semantic",
                    metadata={
                        **metadata,
                        "charCount": len(chunkText),
                        "sentenceCount": len(currentSentences)
                    }
                ))
        
        return chunks
    
    def _fallbackChunk(self, sentences: List[str], documentId: str, documentName: str, metadata: Dict) -> List[DocumentChunk]:
        """回退分块策略（无嵌入服务时）"""
        chunks = []
        currentSentences = []
        currentSize = 0
        chunkIndex = 0
        targetSize = self.targetChunkSize
        
        for sentence in sentences:
            if currentSize + len(sentence) > targetSize and currentSentences:
                chunkText = "".join(currentSentences)
                chunkId = self._generateChunkId(documentId, chunkIndex)
                chunks.append(DocumentChunk(
                    id=chunkId,
                    content=chunkText,
                    documentId=documentId,
                    documentName=documentName,
                    chunkIndex=chunkIndex,
                    chunkType="semantic_fallback",
                    metadata={
                        **metadata,
                        "charCount": len(chunkText),
                        "sentenceCount": len(currentSentences)
                    }
                ))
                chunkIndex += 1
                currentSentences = []
                currentSize = 0
            
            currentSentences.append(sentence)
            currentSize += len(sentence)
        
        if currentSentences:
            chunkText = "".join(currentSentences)
            chunkId = self._generateChunkId(documentId, chunkIndex)
            chunks.append(DocumentChunk(
                id=chunkId,
                content=chunkText,
                documentId=documentId,
                documentName=documentName,
                chunkIndex=chunkIndex,
                chunkType="semantic_fallback",
                metadata={
                    **metadata,
                    "charCount": len(chunkText),
                    "sentenceCount": len(currentSentences)
                }
            ))
        
        return chunks
    
    def _cosineSimilarity(self, vec1: List[float], vec2: List[float]) -> float:
        """计算余弦相似度"""
        import numpy as np
        vec1 = np.array(vec1)
        vec2 = np.array(vec2)
        
        norm1 = np.linalg.norm(vec1)
        norm2 = np.linalg.norm(vec2)
        
        if norm1 == 0 or norm2 == 0:
            return 0.0
        
        return float(np.dot(vec1, vec2) / (norm1 * norm2))
    
    def _generateChunkId(self, documentId: str, chunkIndex: int) -> str:
        """生成分块ID"""
        content = f"{documentId}_{chunkIndex}_{time.time()}"
        return hashlib.md5(content.encode()).hexdigest()[:16]


class HierarchicalChunker(BaseChunker):
    """
    层次化分块器
    
    特点：利用文档结构（标题、章节）创建嵌套分块
    适合：技术文档、论文、手册等结构化文档
    优化：增大最大分块到1200字符，最小分块到200字符
    """
    
    def __init__(self, maxChunkSize: int = 1200, minChunkSize: int = 200):
        self.maxChunkSize = maxChunkSize
        self.minChunkSize = minChunkSize
        
        self.headingPatterns = [
            r'^#{1,6}\s+(.+)$',
            r'^第[一二三四五六七八九十]+[章节部篇]\s*(.*)$',
            r'^\d+\.\s+(.+)$',
            r'^[一二三四五六七八九十]+[、.]\s+(.+)$',
            r'^【(.+)】',
        ]
    
    def chunk(self, text: str, documentId: str, documentName: str, metadata: Optional[Dict] = None) -> List[DocumentChunk]:
        """层次化分块"""
        metadata = metadata or {}
        
        sections = self._extractSections(text)
        
        chunks = []
        chunkIndex = 0
        
        for section in sections:
            sectionChunks = self._processSection(section, documentId, documentName, metadata, chunkIndex)
            chunks.extend(sectionChunks)
            chunkIndex += len(sectionChunks)
        
        return chunks
    
    def _extractSections(self, text: str) -> List[Dict]:
        """提取文档章节"""
        lines = text.split('\n')
        sections = []
        currentSection = {"title": "文档开头", "level": 0, "content": "", "subsections": []}
        
        for line in lines:
            headingLevel, headingTitle = self._detectHeading(line)
            
            if headingLevel > 0:
                if currentSection["content"].strip():
                    sections.append(currentSection)
                currentSection = {
                    "title": headingTitle,
                    "level": headingLevel,
                    "content": "",
                    "subsections": []
                }
            else:
                currentSection["content"] += line + "\n"
        
        if currentSection["content"].strip():
            sections.append(currentSection)
        
        return sections
    
    def _detectHeading(self, line: str) -> tuple:
        """检测标题"""
        line = line.strip()
        if not line:
            return (0, "")
        
        for i, pattern in enumerate(self.headingPatterns):
            match = re.match(pattern, line)
            if match:
                level = i + 1
                title = match.group(1) if match.groups() else line
                return (level, title)
        
        return (0, "")
    
    def _processSection(self, section: Dict, documentId: str, documentName: str, metadata: Dict, startIndex: int) -> List[DocumentChunk]:
        """处理章节"""
        chunks = []
        content = section["content"].strip()
        
        if not content:
            return chunks
        
        if len(content) <= self.maxChunkSize:
            chunkId = self._generateChunkId(documentId, startIndex)
            chunks.append(DocumentChunk(
                id=chunkId,
                content=content,
                documentId=documentId,
                documentName=documentName,
                chunkIndex=startIndex,
                chunkType="hierarchical",
                metadata={
                    **metadata,
                    "charCount": len(content),
                    "sectionTitle": section["title"],
                    "sectionLevel": section["level"]
                }
            ))
        else:
            subChunker = RecursiveChunker(self.maxChunkSize, 50)
            subChunks = subChunker.chunk(content, documentId, documentName, metadata)
            
            for i, subChunk in enumerate(subChunks):
                subChunk.chunkIndex = startIndex + i
                subChunk.chunkType = "hierarchical_sub"
                subChunk.metadata["sectionTitle"] = section["title"]
                subChunk.metadata["sectionLevel"] = section["level"]
                chunks.append(subChunk)
        
        return chunks
    
    def _generateChunkId(self, documentId: str, chunkIndex: int) -> str:
        """生成分块ID"""
        content = f"{documentId}_{chunkIndex}_{time.time()}"
        return hashlib.md5(content.encode()).hexdigest()[:16]


class SmartChunker:
    """
    智能分块器
    
    根据文档类型和内容特征自动选择最佳分块策略
    """
    
    def __init__(self, embeddingService=None):
        self.embeddingService = embeddingService
        self._chunkers = {
            ChunkStrategy.FIXED_SIZE: FixedSizeChunker(),
            ChunkStrategy.RECURSIVE: RecursiveChunker(),
            ChunkStrategy.SEMANTIC: SemanticChunker(embeddingService),
            ChunkStrategy.HIERARCHICAL: HierarchicalChunker(),
        }

    @staticmethod
    def _calculate_overlap(chunk_size: Optional[int]) -> int:
        if not chunk_size:
            return 120
        return max(20, min(int(chunk_size * 0.15), max(chunk_size - 1, 20)))

    def _resolve_chunker(self, strategy: ChunkStrategy, chunkSize: Optional[int]):
        if not chunkSize:
            return self._chunkers.get(strategy)

        normalized_size = max(int(chunkSize), 100)
        overlap = self._calculate_overlap(normalized_size)

        if strategy == ChunkStrategy.FIXED_SIZE:
            return FixedSizeChunker(chunkSize=normalized_size, overlap=overlap)
        if strategy == ChunkStrategy.RECURSIVE:
            return RecursiveChunker(chunkSize=normalized_size, overlap=overlap)
        if strategy == ChunkStrategy.SEMANTIC:
            return SemanticChunker(
                self.embeddingService,
                similarityThreshold=0.6,
                minChunkSize=max(100, int(normalized_size * 0.6)),
                targetChunkSize=normalized_size,
            )
        if strategy == ChunkStrategy.HIERARCHICAL:
            return HierarchicalChunker(
                maxChunkSize=normalized_size,
                minChunkSize=max(100, min(int(normalized_size * 0.4), 400)),
            )
        return self._chunkers.get(strategy)
    
    def chunk(
        self,
        text: str,
        documentId: str,
        documentName: str,
        metadata: Optional[Dict] = None,
        strategy: Optional[ChunkStrategy] = None,
        chunkSize: Optional[int] = None
    ) -> List[DocumentChunk]:
        """
        智能分块
        
        Args:
            text: 文本内容
            documentId: 文档ID
            documentName: 文档名称
            metadata: 元数据
            strategy: 指定策略（可选）
            
        Returns:
            分块列表
        """
        metadata = metadata or {}
        
        if strategy is None:
            strategy = self._detectBestStrategy(text, metadata)

        chunker = self._resolve_chunker(strategy, chunkSize)
        if not chunker:
            chunker = self._chunkers[ChunkStrategy.RECURSIVE]

        size_suffix = f", chunkSize={chunkSize}" if chunkSize else ""
        logger.info(f"使用 {strategy.value} 分块策略处理文档: {documentName}{size_suffix}")
        return chunker.chunk(text, documentId, documentName, metadata)
    
    def _detectBestStrategy(self, text: str, metadata: Dict) -> ChunkStrategy:
        """
        检测最佳分块策略
        
        规则：
        1. 技术文档/论文 -> 层次化分块
        2. 长文本/文章 -> 语义分块
        3. 结构简单文档 -> 递归分块
        4. 日志/数据 -> 固定大小分块
        """
        fileType = metadata.get("fileType", "")
        
        if fileType in [".md", ".rst"]:
            hasHeadings = self._hasMarkdownHeadings(text)
            if hasHeadings:
                return ChunkStrategy.HIERARCHICAL
        
        if fileType in [".pdf", ".docx"]:
            hasStructure = self._hasDocumentStructure(text)
            if hasStructure:
                return ChunkStrategy.HIERARCHICAL
        
        if len(text) > 5000 and self.embeddingService:
            return ChunkStrategy.SEMANTIC
        
        if fileType in [".log", ".csv", ".json"]:
            return ChunkStrategy.FIXED_SIZE
        
        return ChunkStrategy.RECURSIVE
    
    def _hasMarkdownHeadings(self, text: str) -> bool:
        """检测是否有Markdown标题"""
        pattern = r'^#{1,6}\s+'
        return bool(re.search(pattern, text, re.MULTILINE))
    
    def _hasDocumentStructure(self, text: str) -> bool:
        """检测是否有文档结构"""
        patterns = [
            r'^第[一二三四五六七八九十]+[章节部篇]',
            r'^\d+\.\s+[^\d]',
            r'^[一二三四五六七八九十]+[、.]',
        ]
        
        for pattern in patterns:
            if re.search(pattern, text, re.MULTILINE):
                return True
        
        return False


class DocumentParser:
    """
    文档解析器
    
    支持多种格式的文档解析
    """
    
    def __init__(self, ocrExtractor: Optional[Callable[[str], str]] = None):
        self.supportedTypes = {
            ".pdf": self._parsePdf,
            ".docx": self._parseDocx,
            ".doc": self._parseDoc,
            ".xlsx": self._parseXlsx,
            ".xls": self._parseXlsx,
            ".txt": self._parseTxt,
            ".md": self._parseMarkdown,
            ".csv": self._parseCsv,
            ".json": self._parseJson
        }
        self.ocrExtractor = ocrExtractor

    @staticmethod
    def _normalize_scalar_value(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float):
            if value != value:
                return ""
            if value.is_integer():
                return str(int(value))
        text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
        text = re.sub(r"\s+", " ", text)
        return text

    def _format_table_rows(self, rows: List[List[Any]], section_name: str = "") -> str:
        normalized_rows: List[List[str]] = []
        for row in rows:
            normalized = [self._normalize_scalar_value(cell) for cell in list(row or [])]
            if any(cell for cell in normalized):
                normalized_rows.append(normalized)
        if not normalized_rows:
            return ""

        first_row = normalized_rows[0]
        non_empty_headers = [value for value in first_row if value]
        has_header = (
            len(non_empty_headers) >= max(1, len(first_row) // 2)
            and len(non_empty_headers) == len(set(non_empty_headers))
        )
        headers = first_row if has_header else [f"字段{i + 1}" for i in range(len(first_row))]
        data_rows = normalized_rows[1:] if has_header else normalized_rows

        title = f"【表格: {section_name}】" if section_name else "【表格】"
        parts = [title, f"字段：{' | '.join(headers)}"]
        for row_index, row in enumerate(data_rows, start=1):
            pairs = []
            for column_index, value in enumerate(row):
                if not value:
                    continue
                header = headers[column_index] if column_index < len(headers) else f"字段{column_index + 1}"
                pairs.append(f"{header}={value}")
            if pairs:
                parts.append(f"记录{row_index}：{'；'.join(pairs)}")
        return "\n".join(parts).strip()

    def _format_dataframe(self, dataframe: Any, section_name: str = "") -> str:
        try:
            normalized_df = dataframe.fillna("")
        except Exception:
            normalized_df = dataframe

        headers = [self._normalize_scalar_value(column) or f"字段{i + 1}" for i, column in enumerate(list(normalized_df.columns))]
        title = f"【工作表: {section_name}】" if section_name else "【表格数据】"
        parts = [title, f"字段：{' | '.join(headers)}"]

        try:
            rows = normalized_df.itertuples(index=False, name=None)
        except Exception:
            rows = []

        for row_index, row in enumerate(rows, start=1):
            pairs = []
            for column_index, value in enumerate(row):
                normalized_value = self._normalize_scalar_value(value)
                if not normalized_value:
                    continue
                header = headers[column_index] if column_index < len(headers) else f"字段{column_index + 1}"
                pairs.append(f"{header}={normalized_value}")
            if pairs:
                parts.append(f"记录{row_index}：{'；'.join(pairs)}")
        return "\n".join(parts).strip()
    
    def parse(self, filePath: str) -> str:
        """解析文档"""
        ext = os.path.splitext(filePath)[1].lower()
        
        if ext not in self.supportedTypes:
            raise ValueError(f"不支持的文件类型: {ext}")
        
        return self.supportedTypes[ext](filePath)
    
    def _parsePdf(self, filePath: str) -> str:
        """解析PDF文件"""
        try:
            import pypdf
            
            text = []
            with open(filePath, "rb") as f:
                reader = pypdf.PdfReader(f)
                for page in reader.pages:
                    pageText = page.extract_text()
                    if pageText:
                        text.append(pageText)
            if not text:
                if callable(self.ocrExtractor):
                    ocr_text = str(self.ocrExtractor(filePath) or "").strip()
                    if ocr_text:
                        return ocr_text
                raise OCRRequiredError("PDF 未提取到可用文本，可能是扫描件或图片版文档；请先进行 OCR 后再上传")
            return "\n".join(text)
        except ImportError:
            raise ImportError("请安装pypdf: pip install pypdf")
    
    def _parseDocx(self, filePath: str) -> str:
        """解析Word文档"""
        try:
            from docx import Document
            
            doc = Document(filePath)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            
            tables = []
            for table_index, table in enumerate(doc.tables, start=1):
                rows = []
                for row in table.rows:
                    rows.append([cell.text for cell in row.cells])
                formatted = self._format_table_rows(rows, section_name=f"Word表格{table_index}")
                if formatted:
                    tables.append(formatted)
            
            return "\n\n".join(paragraphs + tables)
        except ImportError:
            raise ImportError("请安装python-docx: pip install python-docx")
    
    def _parseDoc(self, filePath: str) -> str:
        """解析旧版Word文档"""
        try:
            import textract
            return textract.process(filePath).decode("utf-8")
        except ImportError:
            return self._parseTxt(filePath)
    
    def _parseXlsx(self, filePath: str) -> str:
        """解析Excel文件"""
        try:
            import pandas as pd
            
            df = pd.read_excel(filePath, sheet_name=None)
            textParts = []
            
            for sheetName, sheetDf in df.items():
                formatted = self._format_dataframe(sheetDf, section_name=str(sheetName or "未命名工作表"))
                if formatted:
                    textParts.append(formatted)
            
            return "\n".join(textParts)
        except ImportError:
            raise ImportError("请安装pandas和openpyxl: pip install pandas openpyxl")
    
    def _parseTxt(self, filePath: str) -> str:
        """解析文本文件"""
        encodings = ['utf-8', 'gbk', 'gb2312', 'utf-16']
        
        for encoding in encodings:
            try:
                with open(filePath, "r", encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        
        with open(filePath, "rb") as f:
            return f.read().decode('utf-8', errors='ignore')

    def _parseMarkdown(self, filePath: str) -> str:
        text = self._parseTxt(filePath)
        text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", lambda m: f"图片说明：{m.group(1).strip()}" if m.group(1).strip() else "", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"```([^\n]*)\n([\s\S]*?)```", lambda m: f"代码块({m.group(1).strip() or 'plain'})：\n{m.group(2).strip()}", text)
        text = self._normalize_markdown_tables(text)
        text = self._normalize_markdown_lists(text)
        return text

    def _normalize_markdown_tables(self, text: str) -> str:
        lines = text.splitlines()
        output: List[str] = []
        index = 0
        table_counter = 1
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            next_line = lines[index + 1].strip() if index + 1 < len(lines) else ""
            if "|" in stripped and re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", next_line):
                table_lines = [stripped]
                index += 2
                while index < len(lines) and "|" in lines[index]:
                    candidate = lines[index].strip()
                    if not candidate:
                        break
                    table_lines.append(candidate)
                    index += 1
                rows: List[List[str]] = []
                for table_line in table_lines:
                    cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                    if any(cells):
                        rows.append(cells)
                formatted = self._format_table_rows(rows, section_name=f"Markdown表格{table_counter}")
                if formatted:
                    output.append(formatted)
                    table_counter += 1
                continue
            output.append(line)
            index += 1
        return "\n".join(output)

    def _normalize_markdown_lists(self, text: str) -> str:
        normalized_lines: List[str] = []
        for raw_line in text.splitlines():
            line = raw_line.rstrip()
            bullet_match = re.match(r"^(\s*)([-*+])\s+(.+)$", line)
            ordered_match = re.match(r"^(\s*)(\d+)\.\s+(.+)$", line)
            if bullet_match:
                indent = len(bullet_match.group(1) or "")
                level = max(indent // 2 + 1, 1)
                normalized_lines.append(f"列表项{level}：{bullet_match.group(3).strip()}")
                continue
            if ordered_match:
                indent = len(ordered_match.group(1) or "")
                level = max(indent // 2 + 1, 1)
                normalized_lines.append(
                    f"有序项{level}.{ordered_match.group(2)}：{ordered_match.group(3).strip()}"
                )
                continue
            normalized_lines.append(line)
        return "\n".join(normalized_lines)

    def _flatten_json_value(self, value: Any, prefix: str = "") -> List[str]:
        lines: List[str] = []
        if isinstance(value, dict):
            for key, item in value.items():
                key_text = self._normalize_scalar_value(key) or "字段"
                next_prefix = f"{prefix}.{key_text}" if prefix else key_text
                lines.extend(self._flatten_json_value(item, next_prefix))
            return lines
        if isinstance(value, list):
            if value and all(isinstance(item, dict) for item in value):
                section_name = prefix or "JSON数组"
                rows = []
                all_headers: List[str] = []
                for item in value:
                    normalized_item = {self._normalize_scalar_value(k) or "字段": self._normalize_scalar_value(v) for k, v in item.items()}
                    rows.append(normalized_item)
                    for key in normalized_item.keys():
                        if key not in all_headers:
                            all_headers.append(key)
                if all_headers:
                    table_rows = [all_headers]
                    for row in rows:
                        table_rows.append([row.get(header, "") for header in all_headers])
                    formatted = self._format_table_rows(table_rows, section_name=section_name)
                    if formatted:
                        lines.append(formatted)
                return lines
            normalized_items = [self._normalize_scalar_value(item) for item in value if self._normalize_scalar_value(item)]
            if normalized_items:
                label = prefix or "列表"
                lines.append(f"{label}：{'；'.join(normalized_items)}")
            return lines
        normalized_value = self._normalize_scalar_value(value)
        if normalized_value:
            label = prefix or "值"
            lines.append(f"{label}：{normalized_value}")
        return lines

    def _parseJson(self, filePath: str) -> str:
        try:
            with open(filePath, "r", encoding="utf-8") as file_obj:
                payload = json.load(file_obj)
        except UnicodeDecodeError:
            with open(filePath, "r", encoding="gbk") as file_obj:
                payload = json.load(file_obj)
        lines = self._flatten_json_value(payload)
        return "\n".join(line for line in lines if line).strip()
    
    def _parseCsv(self, filePath: str) -> str:
        """解析CSV文件"""
        try:
            import pandas as pd
            
            df = pd.read_csv(filePath)
            return self._format_dataframe(df, section_name=os.path.basename(filePath))
        except ImportError:
            return self._parseTxt(filePath)


class IngestionPipeline:
    """
    文档摄取管道
    
    企业级特性：
    - 异步处理队列
    - 进度跟踪
    - 错误重试
    - 文件去重
    - 批量并发
    - 缓存机制
    """
    
    def __init__(
        self,
        embeddingService=None,
        vectorStore=None,
        knowledgeService=None,
        domainProfileService=None,
        knowledgeUnderstandingService=None,
        chunkStrategy: Optional[ChunkStrategy] = None,
        maxRetries: int = 3,
        concurrency: int = 4,
        bm25Retriever=None,
        ocrExtractor: Optional[Callable[[str], str]] = None,
    ):
        self.embeddingService = embeddingService
        self.vectorStore = vectorStore
        self.knowledgeService = knowledgeService
        self.domainProfileService = domainProfileService or get_domain_profile_service()
        self.knowledgeUnderstandingService = knowledgeUnderstandingService or get_knowledge_understanding_service(
            domain_profile_service=self.domainProfileService
        )
        self.chunkStrategy = chunkStrategy
        self.bm25Retriever = bm25Retriever
        
        self.maxRetries = maxRetries
        self.concurrency = concurrency
        self.ocrExtractor = ocrExtractor
        
        self.parser = DocumentParser(ocrExtractor=ocrExtractor)
        self.chunker = SmartChunker(embeddingService)
        
        self._taskQueue = asyncio.Queue()
        self._progress: Dict[str, ProcessingProgress] = {}
        self._fileHashes: Dict[str, str] = {}
        self._documentLedger: Dict[str, Dict[str, Any]] = {}
        self._cache: Dict[str, Any] = {}
        
        self._running = False
        self._workers: List[asyncio.Task] = []
        
        self._loadCache()

    @property
    def _cache_path(self) -> str:
        return str(get_data_dir() / "ingestion_cache.json")

    @staticmethod
    def _atomic_write_json(file_path: str, payload: Dict[str, Any]) -> None:
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        fd, temp_path = tempfile.mkstemp(
            prefix="ingestion_cache_",
            suffix=".tmp",
            dir=os.path.dirname(file_path),
        )
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as file_obj:
                json.dump(payload, file_obj, ensure_ascii=False, indent=2)
            os.replace(temp_path, file_path)
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
    
    def _loadCache(self):
        """加载缓存"""
        cachePath = self._cache_path
        if os.path.exists(cachePath):
            try:
                with open(cachePath, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self._fileHashes = data.get("fileHashes", {})
                    self._documentLedger = data.get("documentLedger", {})
                    self._cache = data.get("transformCache", {})
                logger.info(f"加载缓存: {len(self._fileHashes)} 个文件哈希")
            except Exception as e:
                logger.warning(f"加载缓存失败: {e}")
    
    def _saveCache(self):
        """保存缓存"""
        cachePath = self._cache_path
        try:
            self._atomic_write_json(
                cachePath,
                {
                    "fileHashes": self._fileHashes,
                    "documentLedger": self._documentLedger,
                    "transformCache": self._cache,
                },
            )
        except Exception as e:
            logger.warning(f"保存缓存失败: {e}")

    def set_ocr_extractor(self, ocrExtractor: Optional[Callable[[str], str]]) -> None:
        self.ocrExtractor = ocrExtractor
        self.parser.ocrExtractor = ocrExtractor
    
    def calculateFileHash(self, filePath: str) -> str:
        """计算文件哈希"""
        hasher = hashlib.sha256()
        with open(filePath, "rb") as f:
            for chunk in iter(lambda: f.read(8192), b""):
                hasher.update(chunk)
        return hasher.hexdigest()

    @staticmethod
    def _build_cache_key(enterpriseId: str, contentHash: str) -> str:
        return f"{enterpriseId}_{contentHash}"

    @staticmethod
    def _require_task_enterprise_id(task: Dict[str, Any]) -> str:
        enterprise_id = str((task or {}).get("enterpriseId", "") or "").strip()
        if not enterprise_id:
            raise ValueError("enterpriseId is required for enterprise ingestion pipeline")
        return enterprise_id

    def _build_document_id(self, task: Dict[str, Any]) -> str:
        metadata = task.get("metadata", {}) or {}
        logical_file_id = (
            str(metadata.get("fileId") or "").strip()
            or os.path.normcase(str(task.get("filePath") or "").strip())
            or str(metadata.get("originalName") or "").strip()
        )
        seed = f"{self._require_task_enterprise_id(task)}::{logical_file_id}"
        return hashlib.md5(seed.encode("utf-8")).hexdigest()

    @staticmethod
    def _build_chunk_id(documentId: str, chunkIndex: int, content: str) -> str:
        normalized = str(content or "").strip()
        content_hash = hashlib.md5(normalized.encode("utf-8")).hexdigest()[:12]
        return hashlib.md5(f"{documentId}:{chunkIndex}:{content_hash}".encode("utf-8")).hexdigest()[:16]

    def _stabilize_chunk_ids(self, chunks: List[DocumentChunk]) -> List[DocumentChunk]:
        for index, chunk in enumerate(chunks):
            chunk.chunkIndex = index
            chunk.id = self._build_chunk_id(chunk.documentId, index, chunk.content)
        return chunks

    @staticmethod
    def _utf8_length(text: str) -> int:
        return len(str(text or "").encode("utf-8"))

    def _precleanParsedText(self, text: str, metadata: Optional[Dict[str, Any]] = None) -> str:
        metadata = metadata or {}
        file_type = str(metadata.get("fileType") or "").lower()
        cleaned = str(text or "").replace("\ufeff", "")
        cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
        cleaned = re.sub(r"[ \u3000]+", " ", cleaned)
        cleaned = re.sub(r"\n[ \t]+", "\n", cleaned)
        cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
        cleaned = re.sub(r"(?m)^第\s*\d+\s*页(?:\s*/\s*共\s*\d+\s*页)?\s*$", "", cleaned)
        cleaned = re.sub(r"(?m)^\s*[—_=]{4,}\s*$", "", cleaned)

        deduped_lines: List[str] = []
        previous_line = ""
        for raw_line in cleaned.split("\n"):
            line = raw_line.strip()
            if line == previous_line and line and len(line) <= 80 and file_type in {".pdf", ".docx", ".doc"}:
                continue
            deduped_lines.append(line)
            previous_line = line or previous_line

        cleaned = "\n".join(deduped_lines)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
        return cleaned

    def _extractExplicitQABlocks(self, content: str) -> List[Dict[str, str]]:
        return [
            {
                "question": block["question"],
                "answer": block["answer"],
                "content": block["content"],
            }
            for block in self._extractExplicitQABlocksWithSpans(content)
        ]

    def _extractExplicitQABlocksWithSpans(self, content: str) -> List[Dict[str, Any]]:
        blocks: List[Dict[str, Any]] = []
        seen_pairs: set[tuple[str, str]] = set()

        def append_block(question: str, answer: str, start: int, end: int) -> None:
            normalized_question = str(question or "").strip()
            normalized_answer = str(answer or "").strip()
            normalized_pair = (normalized_question, normalized_answer)
            if normalized_pair in seen_pairs:
                return
            if not self._isValidQA(normalized_question, normalized_answer):
                return
            seen_pairs.add(normalized_pair)
            blocks.append(
                {
                    "question": normalized_question,
                    "answer": normalized_answer,
                    "content": f"问：{normalized_question}\n答：{normalized_answer}",
                    "start": max(int(start or 0), 0),
                    "end": max(int(end or 0), 0),
                }
            )

        pattern_specs = [
            r"Q[：:]\s*(.+?)\s*A[：:]\s*(.+?)(?=Q[：:]|$)",
            r"问题[：:]\s*(.+?)\s*答案[：:]\s*(.+?)(?=问题[：:]|$)",
            r"问[：:]\s*(.+?)\s*答[：:]\s*(.+?)(?=问[：:]|$)",
        ]
        for pattern in pattern_specs:
            for match in re.finditer(pattern, content or "", re.DOTALL):
                append_block(match.group(1), match.group(2), match.start(), match.end())

        if blocks:
            return sorted(blocks, key=lambda item: item["start"])

        question_prefixes = ("问：", "问:", "问题：", "问题:", "Q：", "Q:")
        answer_prefixes = ("答：", "答:", "答案：", "答案:", "A：", "A:")

        def strip_prefix(line: str, prefixes: tuple[str, ...]) -> Optional[str]:
            for prefix in prefixes:
                if line.startswith(prefix):
                    return line[len(prefix):].strip()
            return None

        lines = str(content or "").splitlines(keepends=True)
        current_question: Optional[str] = None
        current_answer: List[str] = []
        block_start: Optional[int] = None
        cursor = 0
        for raw_line in lines:
            line = raw_line.strip()
            line_start = cursor
            cursor += len(raw_line)
            if not line:
                continue

            question_text = strip_prefix(line, question_prefixes)
            if question_text is not None:
                if current_question and current_answer and block_start is not None:
                    append_block(current_question, "\n".join(current_answer), block_start, line_start)
                current_question = question_text
                current_answer = []
                block_start = line_start
                continue

            answer_text = strip_prefix(line, answer_prefixes)
            if answer_text is not None:
                if current_question:
                    current_answer.append(answer_text)
                continue

            if current_question:
                current_answer.append(line)

        if current_question and current_answer and block_start is not None:
            append_block(current_question, "\n".join(current_answer), block_start, cursor)

        return sorted(blocks, key=lambda item: item["start"])

    def _splitLongQaAnswer(self, answer: str, budget_bytes: int) -> List[str]:
        normalized_answer = str(answer or "").strip()
        if not normalized_answer:
            return []

        fragments = [
            fragment.strip()
            for fragment in re.split(r"\n{2,}|(?<=[。！？!?；;])", normalized_answer)
            if fragment.strip()
        ]
        if not fragments:
            fragments = [normalized_answer]

        parts: List[str] = []
        current: List[str] = []
        current_bytes = 0
        for fragment in fragments:
            fragment_bytes = self._utf8_length(fragment)
            if current and current_bytes + fragment_bytes > budget_bytes:
                parts.append("".join(current).strip())
                current = []
                current_bytes = 0

            if fragment_bytes > budget_bytes:
                rough_chars = max(int(budget_bytes / 2), 80)
                for start in range(0, len(fragment), rough_chars):
                    piece = fragment[start:start + rough_chars].strip()
                    if piece:
                        if current:
                            parts.append("".join(current).strip())
                            current = []
                            current_bytes = 0
                        parts.append(piece)
                continue

            current.append(fragment)
            current_bytes += fragment_bytes

        if current:
            parts.append("".join(current).strip())
        return [part for part in parts if part]

    def _buildQaAwareChunks(
        self,
        text: str,
        documentId: str,
        documentName: str,
        metadata: Optional[Dict[str, Any]] = None,
        chunkSize: Optional[int] = None,
    ) -> List[DocumentChunk]:
        metadata = metadata or {}
        qa_blocks = self._extractExplicitQABlocksWithSpans(text)
        category = str(metadata.get("category") or "").lower()
        if not qa_blocks or (len(qa_blocks) < 2 and category != "faq"):
            return []

        chunks: List[DocumentChunk] = []
        chunk_index = 0
        residual_chunker = RecursiveChunker(chunkSize=max(int(chunkSize or 800), 100), overlap=0)
        cursor = 0

        def append_residual_chunks(raw_text: str) -> None:
            nonlocal chunk_index
            residual_text = str(raw_text or "").strip()
            if len(residual_text) < 50:
                return
            residual_chunks = residual_chunker.chunk(
                residual_text,
                documentId,
                documentName,
                {
                    **metadata,
                    "mixedWithQa": True,
                    "residualText": True,
                },
            )
            for residual_chunk in residual_chunks:
                residual_chunk.chunkIndex = chunk_index
                chunk_index += 1
                chunks.append(residual_chunk)

        for qa_index, block in enumerate(qa_blocks):
            append_residual_chunks(text[cursor:block["start"]])
            question = block["question"]
            answer = block["answer"]
            full_text = block["content"]
            qa_group_id = hashlib.md5(f"{documentId}:qa:{qa_index}:{question}".encode("utf-8")).hexdigest()[:16]
            base_metadata = {
                **metadata,
                "qaPairIndex": qa_index,
                "qaGroupId": qa_group_id,
                "questionText": question,
            }
            full_bytes = self._utf8_length(full_text)
            if full_bytes <= QA_SINGLE_CHUNK_MAX_BYTES:
                chunks.append(
                    DocumentChunk(
                        id="",
                        content=full_text,
                        documentId=documentId,
                        documentName=documentName,
                        chunkIndex=chunk_index,
                        chunkType="qa_pair",
                        metadata={
                            **base_metadata,
                            "qaBytes": full_bytes,
                            "qaChunkMode": "single",
                        },
                    )
                )
                chunk_index += 1
                continue

            question_prefix = f"问：{question}\n答："
            answer_budget = max(QA_SINGLE_CHUNK_MAX_BYTES - self._utf8_length(question_prefix), 120)
            answer_parts = self._splitLongQaAnswer(answer, answer_budget)
            for part_index, answer_part in enumerate(answer_parts):
                chunk_content = f"{question_prefix}{answer_part}"
                chunks.append(
                    DocumentChunk(
                        id="",
                        content=chunk_content,
                        documentId=documentId,
                        documentName=documentName,
                        chunkIndex=chunk_index,
                        chunkType="qa_pair_split",
                        metadata={
                            **base_metadata,
                            "qaBytes": self._utf8_length(chunk_content),
                            "qaChunkMode": "split",
                            "qaPartIndex": part_index,
                        },
                    )
                )
                chunk_index += 1

            cursor = block["end"]

        append_residual_chunks(text[cursor:])
        return chunks

    def _build_ledger_key(
        self,
        enterpriseId: str,
        *,
        fileId: str = "",
        documentId: str = "",
        filePath: str = "",
    ) -> str:
        identity = str(fileId or "").strip() or str(documentId or "").strip()
        if not identity:
            identity = os.path.normcase(str(filePath or "").strip())
        return f"{enterpriseId}::{identity}"

    def getDocumentLedger(
        self,
        enterpriseId: str,
        *,
        fileId: str = "",
        documentId: str = "",
        contentHash: str = "",
    ) -> Optional[Dict[str, Any]]:
        keys = [
            self._build_ledger_key(enterpriseId, fileId=fileId),
            self._build_ledger_key(enterpriseId, documentId=documentId),
        ]
        for key in keys:
            if key.endswith("::"):
                continue
            entry = self._documentLedger.get(key)
            if entry:
                return dict(entry)
        for entry in self._documentLedger.values():
            if str(entry.get("enterpriseId") or "") != str(enterpriseId or ""):
                continue
            if fileId and str(entry.get("fileId") or "") == str(fileId):
                return dict(entry)
            if documentId and str(entry.get("documentId") or "") == str(documentId):
                return dict(entry)
            if contentHash and str(entry.get("contentHash") or "") == str(contentHash):
                return dict(entry)
        return None

    def removeDocumentLedger(
        self,
        enterpriseId: str,
        *,
        fileId: str = "",
        documentId: str = "",
        contentHash: str = "",
        deleteStage: str = "deleted",
    ) -> Optional[Dict[str, Any]]:
        removed_entry: Optional[Dict[str, Any]] = None
        keys_to_update: List[str] = []
        for key, entry in self._documentLedger.items():
            if str(entry.get("enterpriseId") or "") != str(enterpriseId or ""):
                continue
            if fileId and str(entry.get("fileId") or "") == str(fileId):
                keys_to_update.append(key)
            elif documentId and str(entry.get("documentId") or "") == str(documentId):
                keys_to_update.append(key)
            elif contentHash and str(entry.get("contentHash") or "") == str(contentHash):
                keys_to_update.append(key)
        for key in keys_to_update:
            current_entry = dict(self._documentLedger.get(key) or {})
            removed_entry = current_entry or removed_entry
            if not current_entry:
                continue
            current_entry["status"] = "deleted"
            current_entry["deletedAt"] = datetime.now().isoformat()
            current_entry["deleteStage"] = str(deleteStage or "deleted")
            self._documentLedger[key] = current_entry
        if keys_to_update:
            removed_hashes = {
                str((removed_entry or {}).get("contentHash") or ""),
                str(contentHash or ""),
            }
            for ledger_hash in removed_hashes:
                if ledger_hash:
                    self._fileHashes.pop(self._build_cache_key(enterpriseId, ledger_hash), None)
            self._saveCache()
        return dict(removed_entry) if removed_entry else None
    
    def isDuplicate(self, filePath: str, enterpriseId: str) -> bool:
        """检查文件是否重复"""
        fileHash = self.calculateFileHash(filePath)
        cacheKey = self._build_cache_key(enterpriseId, fileHash)
        return cacheKey in self._fileHashes

    @staticmethod
    def _estimateSentenceCount(text: str) -> int:
        parts = [part.strip() for part in re.split(r"[。！？!?；;\n]+", text or "") if part.strip()]
        return len(parts)

    def _buildChunkQuality(self, chunk: DocumentChunk) -> Dict[str, Any]:
        content = chunk.content or ""
        char_count = len(content)
        word_count = len(content.split()) if content else 0
        sentence_count = self._estimateSentenceCount(content)
        overlap_chars = int(chunk.metadata.get("overlapChars", 0) or 0)
        overlap_ratio = round(min(overlap_chars / max(char_count, 1), 1.0), 3) if char_count else 0.0
        visible_chars = [ch for ch in content if not ch.isspace()]
        signal_chars = re.findall(r"[\u4e00-\u9fffA-Za-z0-9]", content)
        text_signal_ratio = round(len(signal_chars) / max(len(visible_chars), 1), 3) if visible_chars else 0.0
        normalized_compact = re.sub(r"\s+", "", content)
        repeated_char_ratio = 0.0
        if normalized_compact:
            repeated_char_ratio = round(
                max(normalized_compact.count(ch) for ch in set(normalized_compact)) / len(normalized_compact),
                3,
            )

        size_score = min(char_count / 400, 1.0) if char_count < 400 else max(0.0, 1.0 - ((char_count - 900) / 900))
        sentence_score = min(sentence_count / 3, 1.0)
        overlap_score = 1.0 - min(abs(overlap_ratio - 0.15) / 0.15, 1.0) if overlap_ratio > 0 else 0.6
        quality_score = round(max(0.0, min(1.0, size_score * 0.5 + sentence_score * 0.3 + overlap_score * 0.2)), 3)
        is_noise_only = bool(normalized_compact) and bool(re.fullmatch(r"[-—_=|•·*#.`~]+", normalized_compact))
        review_required = bool(
            is_noise_only
            or (char_count >= 20 and text_signal_ratio < 0.25 and not str(chunk.chunkType or "").startswith("qa_pair"))
            or (len(normalized_compact) >= 30 and repeated_char_ratio >= 0.75)
        )

        return {
            "chunk_id": chunk.id,
            "document_id": chunk.documentId,
            "chunk_type": chunk.chunkType,
            "char_count": char_count,
            "word_count": word_count,
            "sentence_count": sentence_count,
            "overlap_chars": overlap_chars,
            "overlap_ratio": overlap_ratio,
            "quality_score": quality_score,
            "text_signal_ratio": text_signal_ratio,
            "repeated_char_ratio": repeated_char_ratio,
            "review_required": review_required,
        }

    def _annotateChunkQuality(self, chunks: List[DocumentChunk]) -> Dict[str, Any]:
        if not chunks:
            return {
                "total_chunks": 0,
                "average_char_count": 0.0,
                "average_sentence_count": 0.0,
                "average_quality_score": 0.0,
                "small_chunk_ratio": 0.0,
                "large_chunk_ratio": 0.0,
                "chunk_type_distribution": {},
            }

        type_distribution: Dict[str, int] = {}
        total_chars = 0
        total_sentences = 0
        total_quality = 0.0
        small_chunks = 0
        large_chunks = 0

        for chunk in chunks:
            quality = self._buildChunkQuality(chunk)
            chunk.metadata["chunkQuality"] = quality
            type_distribution[quality["chunk_type"]] = type_distribution.get(quality["chunk_type"], 0) + 1
            total_chars += quality["char_count"]
            total_sentences += quality["sentence_count"]
            total_quality += quality["quality_score"]
            if quality["char_count"] < 120:
                small_chunks += 1
            if quality["char_count"] > 1200:
                large_chunks += 1

        chunk_count = len(chunks)
        return {
            "total_chunks": chunk_count,
            "average_char_count": round(total_chars / chunk_count, 1),
            "average_sentence_count": round(total_sentences / chunk_count, 1),
            "average_quality_score": round(total_quality / chunk_count, 3),
            "small_chunk_ratio": round(small_chunks / chunk_count, 3),
            "large_chunk_ratio": round(large_chunks / chunk_count, 3),
            "chunk_type_distribution": type_distribution,
        }

    def _filterLowSignalChunks(self, chunks: List[DocumentChunk]) -> tuple[List[DocumentChunk], int]:
        filtered: List[DocumentChunk] = []
        removed_count = 0
        for chunk in chunks:
            quality = (chunk.metadata or {}).get("chunkQuality") or self._buildChunkQuality(chunk)
            chunk.metadata["chunkQuality"] = quality
            if quality.get("review_required") and not str(chunk.chunkType or "").startswith("qa_pair"):
                removed_count += 1
                continue
            filtered.append(chunk)
        if filtered:
            return filtered, removed_count
        return chunks, 0
    
    async def submitTask(
        self,
        filePath: str,
        enterpriseId: str,
        category: str = "general",
        tags: Optional[List[str]] = None,
        metadata: Optional[Dict] = None,
        chunkStrategy: Optional[ChunkStrategy] = None,
        chunkSize: Optional[int] = None
    ) -> str:
        """
        提交处理任务
        
        Args:
            filePath: 文件路径
            enterpriseId: 企业ID
            category: 分类
            tags: 标签
            metadata: 元数据
            
        Returns:
            任务ID
        """
        taskId = str(uuid.uuid4())
        
        if self.isDuplicate(filePath, enterpriseId):
            logger.info(f"文件已存在，跳过: {filePath}")
            return None
        
        task = {
            "taskId": taskId,
            "filePath": filePath,
            "enterpriseId": enterpriseId,
            "category": category,
            "tags": tags or [],
            "metadata": metadata or {},
            "chunkStrategy": chunkStrategy,
            "chunkSize": chunkSize,
        }
        
        progress = ProcessingProgress(
            taskId=taskId,
            fileName=os.path.basename(filePath),
            status=ProcessingStatus.PENDING,
            startTime=datetime.now()
        )
        self._progress[taskId] = progress
        
        await self._taskQueue.put(task)
        
        if not self._running:
            # 使用 await 确保工作线程正确启动
            await self._startWorkers()
        
        return taskId
    
    async def submitBatch(
        self,
        tasks: List[Dict]
    ) -> List[str]:
        """
        批量提交任务
        
        Args:
            tasks: 任务列表
            
        Returns:
            任务ID列表
        """
        taskIds = []
        for task in tasks:
            taskId = await self.submitTask(**task)
            if taskId:
                taskIds.append(taskId)
        return taskIds
    
    async def _startWorkers(self):
        """启动工作线程"""
        if self._running:
            return
        
        self._running = True
        
        for i in range(self.concurrency):
            worker = asyncio.create_task(self._worker(i))
            self._workers.append(worker)
        
        logger.info(f"启动 {self.concurrency} 个处理工作线程")
    
    async def _worker(self, workerId: int):
        """工作线程"""
        while self._running:
            try:
                task = await asyncio.wait_for(self._taskQueue.get(), timeout=1.0)
                await self._processTask(task)
            except asyncio.TimeoutError:
                continue
            except Exception as e:
                logger.error(f"工作线程 {workerId} 错误: {e}")
    
    async def _processTask(self, task: Dict):
        """处理任务"""
        taskId = task["taskId"]
        progress = self._progress.get(taskId)
        
        if not progress:
            return
        
        try:
            progress.status = ProcessingStatus.PARSING
            progress.currentStep = "解析文档"
            progress.progress = 0.1
            
            text = await self._parseWithRetry(task["filePath"])
            
            progress.status = ProcessingStatus.CHUNKING
            progress.currentStep = "文档分块"
            progress.progress = 0.3
            
            contentHash = self.calculateFileHash(task["filePath"])
            documentId = self._build_document_id(task)
            progress.documentId = documentId
            progress.contentHash = contentHash
            task.setdefault("metadata", {})
            task["metadata"]["documentId"] = documentId
            
            docMetadata = {
                **task.get("metadata", {}),
                "enterpriseId": task["enterpriseId"],
                "enterprise_id": task["enterpriseId"],
                "documentId": documentId,
                "documentName": progress.fileName,
                "category": task["category"],
                "tags": task["tags"],
                "fileType": os.path.splitext(task["filePath"])[1].lower(),
            }
            cleaned_text = self._precleanParsedText(text, docMetadata)
            
            chunks = self._buildQaAwareChunks(
                text=cleaned_text,
                documentId=documentId,
                documentName=progress.fileName,
                metadata=docMetadata,
                chunkSize=task.get("chunkSize"),
            )
            if not chunks:
                chunks = self.chunker.chunk(
                    text=cleaned_text,
                    documentId=documentId,
                    documentName=progress.fileName,
                    metadata=docMetadata,
                    strategy=task.get("chunkStrategy") or self.chunkStrategy,
                    chunkSize=task.get("chunkSize"),
                )
            if not chunks and cleaned_text.strip():
                fallback_text = cleaned_text.strip()
                chunks = [
                    DocumentChunk(
                        id=hashlib.md5(f"{documentId}_fallback_0".encode()).hexdigest()[:16],
                        content=fallback_text,
                        documentId=documentId,
                        documentName=progress.fileName,
                        chunkIndex=0,
                        chunkType="fallback_single",
                        metadata={
                            **docMetadata,
                            "fallbackReason": "no_chunks_after_chunking",
                        },
                    )
                ]
            chunks = self._stabilize_chunk_ids(chunks)
            progress.chunkCount = len(chunks)
            progress.chunkIds = [chunk.id for chunk in chunks]
            progress.chunkQualitySummary = self._annotateChunkQuality(chunks)
            chunks, filtered_count = self._filterLowSignalChunks(chunks)
            if filtered_count:
                chunks = self._stabilize_chunk_ids(chunks)
                progress.chunkCount = len(chunks)
                progress.chunkIds = [chunk.id for chunk in chunks]
                progress.chunkQualitySummary = self._annotateChunkQuality(chunks)
            progress.chunkQualitySummary["filtered_chunk_count"] = filtered_count
            
            progress.status = ProcessingStatus.EMBEDDING
            progress.currentStep = "生成向量"
            progress.progress = 0.5
            
            if self.embeddingService:
                await self._embedChunks(chunks)
            
            progress.status = ProcessingStatus.STORING
            progress.currentStep = "存储向量"
            progress.progress = 0.7
            
            if self.vectorStore:
                await self._storeChunks(chunks, task["enterpriseId"])
            
            progress.progress = 0.9
            progress.currentStep = "提取知识点"
            
            if self.knowledgeService:
                await self._extractKnowledge(chunks, task)
            
            progress.status = ProcessingStatus.COMPLETED
            progress.progress = 1.0
            progress.currentStep = "完成"
            progress.endTime = datetime.now()
            progress.vectorCount = len(chunks)
            progress.knowledgeItemIds = list(dict.fromkeys(task.get("_knowledge_item_ids", []) or []))
            ledger_entry = {
                "enterpriseId": task["enterpriseId"],
                "fileId": str((task.get("metadata", {}) or {}).get("fileId") or ""),
                "documentId": documentId,
                "contentHash": contentHash,
                "filePath": task["filePath"],
                "originalName": str((task.get("metadata", {}) or {}).get("originalName") or progress.fileName),
                "category": task.get("category", "general"),
                "tags": list(task.get("tags", []) or []),
                "chunkIds": list(progress.chunkIds),
                "knowledgeItemIds": list(progress.knowledgeItemIds),
                "chunkCount": progress.chunkCount,
                "vectorCount": progress.vectorCount,
                "updatedAt": datetime.now().isoformat(),
                "status": progress.status.value,
            }
            progress.indexLedger = dict(ledger_entry)
            ledger_key = self._build_ledger_key(
                task["enterpriseId"],
                fileId=ledger_entry["fileId"],
                documentId=documentId,
                filePath=task["filePath"],
            )
            self._documentLedger[ledger_key] = ledger_entry
            cacheKey = self._build_cache_key(task["enterpriseId"], contentHash)
            self._fileHashes[cacheKey] = documentId
            self._saveCache()
            
            logger.info(f"任务完成: {progress.fileName}, 分块数: {progress.chunkCount}")
            
        except Exception as e:
            import traceback
            if isinstance(e, OCRRequiredError):
                progress.status = ProcessingStatus.OCR_REQUIRED
                progress.currentStep = "等待OCR处理"
                progress.errorMessage = str(e)
            else:
                progress.status = ProcessingStatus.FAILED
                progress.errorMessage = f"{str(e)}\n{traceback.format_exc()}"
            progress.endTime = datetime.now()
            logger.error(f"任务失败: {progress.fileName}, 错误: {e}")
            logger.error(f"错误堆栈: {traceback.format_exc()}")
    
    async def _parseWithRetry(self, filePath: str) -> str:
        """带重试的解析"""
        lastError = None
        
        for attempt in range(self.maxRetries):
            try:
                return await asyncio.get_event_loop().run_in_executor(
                    None, self.parser.parse, filePath
                )
            except Exception as e:
                lastError = e
                logger.warning(f"解析失败 (尝试 {attempt + 1}/{self.maxRetries}): {e}")
                await asyncio.sleep(1)
        
        raise lastError
    
    async def _embedChunks(self, chunks: List[DocumentChunk]):
        """生成向量嵌入"""
        batchSize = 32
        
        for i in range(0, len(chunks), batchSize):
            batch = chunks[i:i + batchSize]
            contents = [c.content for c in batch]
            
            try:
                embeddings = self.embeddingService.embed(contents)
                for j, chunk in enumerate(batch):
                    chunk.embedding = embeddings[j].tolist() if hasattr(embeddings[j], 'tolist') else embeddings[j]
            except Exception as e:
                logger.warning(f"批量嵌入失败: {e}")
            
            await asyncio.sleep(0.01)
    
    async def _storeChunks(self, chunks: List[DocumentChunk], enterpriseId: str):
        """存储到向量数据库"""
        documents = []
        
        for chunk in chunks:
            metadata = {
                **(chunk.metadata or {}),
                "enterpriseId": enterpriseId,
                "enterprise_id": enterpriseId,
                "documentId": chunk.documentId,
                "documentName": chunk.documentName,
                "chunkId": chunk.id,
                "chunkIndex": chunk.chunkIndex,
                "chunkType": chunk.chunkType,
            }
            doc = {
                "id": chunk.id,
                "content": chunk.content,
                "metadata": metadata,
            }
            if chunk.embedding:
                doc["embedding"] = chunk.embedding
            documents.append(doc)
        
        if documents:
            try:
                self.vectorStore.addDocuments(enterpriseId, documents)
            except Exception as e:
                logger.error(f"存储向量失败: {e}")
                raise

            if self.bm25Retriever is not None:
                try:
                    self.bm25Retriever.add_documents(documents)
                except Exception as e:
                    logger.warning(f"BM25索引同步失败: {e}")
    
    async def _extractKnowledge(self, chunks: List[DocumentChunk], task: Dict):
        """提取知识点并同步到知识库管理系统，自动进行三层分类"""
        extractedCount = 0
        candidateCount = 0
        saveErrors: List[str] = []
        extracted_items: List[Dict[str, Any]] = []
        knowledge_item_ids: List[str] = []
        documentName = task.get("metadata", {}).get("originalName", "未知文档")
        seen_qa: set[tuple[str, str]] = set()
        combined_content = "\n\n".join(
            chunk.content.strip() for chunk in chunks if str(chunk.content or "").strip()
        )

        qa_sources: List[tuple[DocumentChunk, List[tuple[str, str]]]] = []
        if combined_content:
            synthetic_chunk = chunks[0] if chunks else DocumentChunk(
                id="",
                content=combined_content,
                documentId=task.get("metadata", {}).get("documentId", ""),
                documentName=documentName,
                chunkIndex=0,
                metadata={},
            )
            combined_pairs = self._extractQAFromChunk(combined_content)
            if combined_pairs:
                qa_sources.append((synthetic_chunk, combined_pairs))

        if not qa_sources:
            for chunk in chunks:
                if len(chunk.content) < 20:
                    continue
                qaPairs = self._extractQAFromChunk(chunk.content)
                if qaPairs:
                    qa_sources.append((chunk, qaPairs))

        for chunk, qaPairs in qa_sources:
            enterprise_id = self._require_task_enterprise_id(task)
            # 解析当前企业激活的 schema_id，绑定到知识条目避免跨 schema 扫描
            try:
                schema_id = get_industry_schema_service().get_effective_active_schema_id(
                    enterprise_id=enterprise_id
                ) or ""
            except Exception:
                schema_id = ""
            for question, answer in qaPairs:
                normalized_pair = (str(question or "").strip(), str(answer or "").strip())
                if normalized_pair in seen_qa:
                    continue
                seen_qa.add(normalized_pair)
                if len(question) > 5 and len(answer) > 10:
                    candidateCount += 1
                    category = self._smartCategorize(
                        question,
                        answer,
                        task.get("category", "other"),
                        enterprise_id=str(task.get("enterpriseId", "") or ""),
                    )
                    domain, topic, tags = self._classifyHierarchy(question, answer, category)

                    itemData = {
                        "question": question,
                        "answer": answer,
                        "category": category,
                        "domain": domain,
                        "topic": topic,
                        "tags": tags,
                        "keywords": self._extractKeywords(question + " " + answer),
                        "source": "document_upload",
                        "enterprise_id": enterprise_id,
                        "schema_id": schema_id,
                        "enabled": True,
                        "priority": 10,
                        "metadata": {
                            "fileId": task.get("metadata", {}).get("fileId", ""),
                            "documentId": chunk.documentId,
                            "chunkId": chunk.id,
                            "chunkIndex": chunk.chunkIndex,
                            "chunkType": chunk.chunkType,
                            "chunkQuality": (chunk.metadata or {}).get("chunkQuality", {}),
                            "originalName": documentName,
                            "extractedAt": datetime.now().isoformat()
                        }
                    }

                    try:
                        created_item = self.knowledgeService.add_item(itemData)
                        extractedCount += 1
                        if created_item and getattr(created_item, "id", ""):
                            knowledge_item_ids.append(created_item.id)
                        extracted_items.append(itemData)
                    except Exception as e:
                        saveErrors.append(str(e))
                        logger.warning(f"添加知识点失败: {e}")

        task["_knowledge_item_ids"] = list(dict.fromkeys(knowledge_item_ids))
        await self._buildKnowledgeProfile(chunks, task, combined_content, extracted_items)

        if extractedCount > 0:
            logger.info(f"从文档 '{documentName}' 提取并保存了 {extractedCount} 个知识点")
            await self._notifyKnowledgeUpdate(extractedCount)
            return extractedCount
        
        if candidateCount > 0:
            raise RuntimeError(saveErrors[0] if saveErrors else f"文档 '{documentName}' 识别出问答，但知识点保存失败")
        
        default_created = await self._createDefaultKnowledge(chunks, task)
        if not default_created:
            raise RuntimeError(f"文档 '{documentName}' 未能生成可入库知识")
        return 1

    async def _buildKnowledgeProfile(
        self,
        chunks: List[DocumentChunk],
        task: Dict,
        combined_content: str,
        extracted_items: List[Dict[str, Any]],
    ) -> None:
        """基于上传文档生成企业知识画像，为通用意图识别和检索路由提供基础画像。"""
        if not self.knowledgeUnderstandingService:
            return
        content = str(combined_content or "").strip()
        if not content:
            content = "\n\n".join(
                str(chunk.content or "").strip()
                for chunk in chunks[:8]
                if str(chunk.content or "").strip()
            )
        if not content:
            return

        metadata = dict(task.get("metadata", {}) or {})
        metadata.setdefault("category", task.get("category", "other"))
        metadata.setdefault("tags", list(task.get("tags", []) or []))
        document_name = metadata.get("originalName", "未知文档")
        enterprise_id = self._require_task_enterprise_id(task)

        try:
            await self.knowledgeUnderstandingService.build_and_store_profile(
                enterprise_id=enterprise_id,
                document_name=document_name,
                content=content,
                extracted_items=extracted_items,
                metadata=metadata,
            )
        except Exception as exc:
            logger.warning(f"生成知识画像失败: {exc}")
    
    def _classifyHierarchy(self, question: str, answer: str, category: str) -> tuple:
        """
        三层分类：领域→主题→标签
        
        Args:
            question: 问题
            answer: 答案
            category: 基础分类
        
        Returns:
            Tuple[domain, topic, tags]
        """
        try:
            from src.common.graph_taxonomy import CATEGORY_HIERARCHY, CATEGORY_TO_DOMAIN
            text = (question + " " + answer).lower()
            
            domain = CATEGORY_TO_DOMAIN.get(category, "tips")
            
            topic = self._matchTopicInDomain(domain, text)
            if not topic:
                for dId, dData in CATEGORY_HIERARCHY.items():
                    matched = self._matchTopicInDomain(dId, text)
                    if matched:
                        domain = dId
                        topic = matched
                        break
            
            if not topic:
                firstTopic = list(CATEGORY_HIERARCHY.get(domain, {}).get("topics", {}).keys())
                topic = firstTopic[0] if firstTopic else "unknown"
            
            tags = self._extractHierarchyTags(text, domain, topic)
            
            return domain, topic, tags
        except Exception as e:
            logger.warning(f"三层分类失败: {e}")
            return CATEGORY_TO_DOMAIN.get(category, "tips"), "unknown", []
    
    def _matchTopicInDomain(self, domainId: str, text: str) -> str:
        """在指定领域内匹配主题"""
        try:
            from src.common.graph_taxonomy import CATEGORY_HIERARCHY
            domainData = CATEGORY_HIERARCHY.get(domainId, {})
            topics = domainData.get("topics", {})
            
            bestTopic = None
            bestScore = 0
            
            for topicId, topicData in topics.items():
                score = 0
                topic_terms = topicData.get("tags", []) or topicData.get("keywords", [])
                for kw in topic_terms:
                    if kw.lower() in text:
                        score += 1
                if score > bestScore:
                    bestScore = score
                    bestTopic = topicId
            
            return bestTopic
        except Exception:
            return None
    
    def _extractHierarchyTags(self, text: str, domain: str, topic: str) -> List[str]:
        """提取层次化标签"""
        try:
            from src.common.graph_taxonomy import CATEGORY_HIERARCHY
            tags = []
            domainData = CATEGORY_HIERARCHY.get(domain, {})
            domainLabel = domainData.get("label", "")
            if domainLabel:
                tags.append(domainLabel)
            
            topicData = domainData.get("topics", {}).get(topic, {})
            topicLabel = topicData.get("label", "")
            if topicLabel:
                tags.append(topicLabel)
            
            tagPatterns = [
                "智能获客", "营销自动化", "数据采集", "意向分析", "消息群发",
                "客服机器人", "客户管理", "抖音", "小红书", "微信",
                "SaaS", "私有化", "AI", "大数据", "自动化",
                "代理", "加盟", "售后", "培训", "投诉"
            ]
            for p in tagPatterns:
                if p.lower() in text and p not in tags:
                    tags.append(p)
            
            return tags[:8]
        except Exception:
            return []
    
    def _smartCategorize(
        self,
        question: str,
        answer: str,
        defaultCategory: str,
        *,
        enterprise_id: str = "",
    ) -> str:
        """
        智能分类知识点

        优先按当前活动 Schema 的分类关键词进行归类，缺失时回退默认分类。
        """
        text = f"{question} {answer}".lower()
        return infer_category_from_text(
            text,
            default_category=defaultCategory or "other",
            enterprise_id=str(enterprise_id or "").strip(),
        )
    
    async def _notifyKnowledgeUpdate(self, count: int):
        """
        通知知识库更新
        
        Args:
            count: 新增知识点数量
        """
        try:
            # 重置知识库服务缓存，确保下次加载获取最新数据
            if hasattr(self.knowledgeService, '_invalidate_cache'):
                self.knowledgeService._invalidate_cache()
            logger.info(f"知识库已更新，新增 {count} 条知识点")
        except Exception as e:
            logger.warning(f"通知知识库更新失败: {e}")
    
    async def _createDefaultKnowledge(self, chunks: List[DocumentChunk], task: Dict) -> bool:
        """
        创建默认知识点
        
        当文档无法提取结构化知识点时，创建基于文档内容的默认知识点
        """
        documentName = task.get("metadata", {}).get("originalName", "未知文档")
        
        # 合并所有分块内容
        allContent = "\n\n".join([chunk.content for chunk in chunks if len(str(chunk.content or "").strip()) > 0])
        
        if len(allContent) >= 20:
            enterprise_id = self._require_task_enterprise_id(task)
            # 解析当前企业激活的 schema_id，绑定到知识条目避免跨 schema 扫描
            try:
                schema_id = get_industry_schema_service().get_effective_active_schema_id(
                    enterprise_id=enterprise_id
                ) or ""
            except Exception:
                schema_id = ""
            # 创建摘要作为答案
            summary = allContent[:500] + "..." if len(allContent) > 500 else allContent

            itemData = {
                "question": f"关于{documentName}的内容",
                "answer": summary,
                "category": task.get("category", "other"),
                "keywords": self._extractKeywords(allContent),
                "source": "document_upload",
                "enterprise_id": enterprise_id,
                "schema_id": schema_id,
                "enabled": True,
                "priority": 5,
                "metadata": {
                    "fileId": task.get("metadata", {}).get("fileId", ""),
                    "documentId": chunks[0].documentId if chunks else "",
                    "chunkType": chunks[0].chunkType if chunks else "",
                    "originalName": documentName,
                    "isDefault": True,
                    "extractedAt": datetime.now().isoformat()
                }
            }
            
            try:
                created_item = self.knowledgeService.add_item(itemData)
                task.setdefault("_knowledge_item_ids", [])
                if created_item and getattr(created_item, "id", ""):
                    task["_knowledge_item_ids"].append(created_item.id)
                logger.info(f"为文档 '{documentName}' 创建了默认知识点")
                return True
            except Exception as e:
                logger.warning(f"创建默认知识点失败: {e}")
                return False
        return False
    
    def _isValidQA(self, question: str, answer: str) -> bool:
        """
        验证问答对是否有效
        
        过滤无效内容：
        1. 问题以分隔线开头
        2. 问题主要是分隔线字符
        3. 答案太短或无效
        """
        if not question or not answer:
            return False
        
        question = question.strip()
        answer = answer.strip()
        
        if len(question) < 3 or len(answer) < 5:
            return False
        
        separatorChars = '─═_*#`. \n\r\t-—–'
        
        if question[0] in '─═_':
            return False
        
        questionStripped = question.strip(separatorChars)
        if len(questionStripped) < 3:
            return False
        
        separatorCount = sum(question.count(c) for c in '─═_')
        if separatorCount > len(question) * 0.5:
            return False
        
        answerStripped = answer.strip(separatorChars)
        if len(answerStripped) < 5:
            return False
        
        invalidPatterns = [
            r'^[─═_\-—–]+$',
            r'^\.{3,}$',
            r'^\*+$',
            r'^#+$',
        ]
        for pattern in invalidPatterns:
            if re.match(pattern, question):
                return False
            if re.match(pattern, answer):
                return False
        
        return True
    
    def _extractQAFromChunk(self, content: str) -> List[tuple]:
        """从文本块提取问答对"""
        qaPairs = []
        
        qaPattern1 = r'Q[：:]\s*(.+?)\s*A[：:]\s*(.+?)(?=Q[：:]|$)'
        matches1 = re.findall(qaPattern1, content, re.DOTALL)
        for q, a in matches1:
            q, a = q.strip(), a.strip()
            if self._isValidQA(q, a):
                qaPairs.append((q, a))
        
        qaPattern2 = r'问题[：:]\s*(.+?)\s*答案[：:]\s*(.+?)(?=问题[：:]|$)'
        matches2 = re.findall(qaPattern2, content, re.DOTALL)
        for q, a in matches2:
            q, a = q.strip(), a.strip()
            if self._isValidQA(q, a):
                qaPairs.append((q, a))

        qaPattern3 = r'问[：:]\s*(.+?)\s*答[：:]\s*(.+?)(?=问[：:]|$)'
        matches3 = re.findall(qaPattern3, content, re.DOTALL)
        for q, a in matches3:
            q, a = q.strip(), a.strip()
            if self._isValidQA(q, a):
                qaPairs.append((q, a))
        
        if not qaPairs:
            question_prefixes = ("问：", "问:", "问题：", "问题:", "Q：", "Q:")
            answer_prefixes = ("答：", "答:", "答案：", "答案:", "A：", "A:")

            def strip_prefix(line: str, prefixes: tuple) -> Optional[str]:
                for prefix in prefixes:
                    if line.startswith(prefix):
                        return line[len(prefix):].strip()
                return None

            currentQuestion = None
            currentAnswer = []
            for rawLine in content.splitlines():
                line = rawLine.strip()
                if not line:
                    continue

                questionText = strip_prefix(line, question_prefixes)
                if questionText is not None:
                    if currentQuestion and currentAnswer:
                        answer = "\n".join(currentAnswer).strip()
                        if self._isValidQA(currentQuestion, answer):
                            qaPairs.append((currentQuestion, answer))
                    currentQuestion = questionText
                    currentAnswer = []
                    continue

                answerText = strip_prefix(line, answer_prefixes)
                if answerText is not None:
                    if currentQuestion:
                        currentAnswer.append(answerText)
                    continue

                if currentQuestion:
                    currentAnswer.append(line)

            if currentQuestion and currentAnswer:
                answer = "\n".join(currentAnswer).strip()
                if self._isValidQA(currentQuestion, answer):
                    qaPairs.append((currentQuestion, answer))

        if not qaPairs:
            lines = content.split('\n')
            currentTitle = None
            currentContent = []
            
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                    
                if len(line) < 50 and (line.endswith('？') or line.endswith('?') or 
                    line.startswith(('如何', '怎么', '什么是', '为什么', '能不能', '是否'))):
                    if currentTitle and currentContent:
                        answer = '\n'.join(currentContent)
                        if self._isValidQA(currentTitle, answer):
                            qaPairs.append((currentTitle, answer))
                    currentTitle = line
                    currentContent = []
                else:
                    if currentTitle:
                        currentContent.append(line)
                    elif len(line) > 50:
                        question = line[:30] + "..." if len(line) > 30 else line
                        if self._isValidQA(question, line):
                            qaPairs.append((question, line))
            
            if currentTitle and currentContent:
                answer = '\n'.join(currentContent)
                if self._isValidQA(currentTitle, answer):
                    qaPairs.append((currentTitle, answer))
        
        if not qaPairs:
            paragraphs = content.split('\n\n')
            for para in paragraphs:
                para = para.strip()
                if len(para) > 50:
                    question = self._generateQuestionFromContent(para)
                    if self._isValidQA(question, para):
                        qaPairs.append((question, para))
        
        return qaPairs[:10]
    
    def _generateQuestionFromContent(self, content: str) -> str:
        """从内容生成问题"""
        # 提取关键词作为问题基础
        keywords = self._extractKeywords(content)
        if keywords:
            return f"关于{keywords[0]}的说明"
        
        # 使用内容前30字符
        if len(content) > 30:
            return content[:30] + "..."
        return content
    
    def _extractKeywords(self, text: str) -> List[str]:
        """提取关键词"""
        chinesePattern = r'[\u4e00-\u9fa5]{2,4}'
        matches = re.findall(chinesePattern, text)
        
        stopWords = {'的', '是', '在', '有', '和', '了', '对', '这', '那', '我', '你', '他', '她', '它'}
        
        keywords = []
        for word in matches:
            if word not in stopWords and word not in keywords:
                keywords.append(word)
        
        return keywords[:5]
    
    def getProgress(self, taskId: str) -> Optional[ProcessingProgress]:
        """获取任务进度"""
        return self._progress.get(taskId)
    
    def getAllProgress(self) -> List[ProcessingProgress]:
        """获取所有任务进度"""
        return list(self._progress.values())
    
    def cancelTask(self, taskId: str) -> bool:
        """取消任务"""
        progress = self._progress.get(taskId)
        if progress and progress.status in [ProcessingStatus.PENDING, ProcessingStatus.PARSING]:
            progress.status = ProcessingStatus.CANCELLED
            progress.endTime = datetime.now()
            return True
        return False
    
    async def shutdown(self):
        """关闭管道"""
        self._running = False
        
        for worker in self._workers:
            worker.cancel()
        
        self._workers.clear()
        self._saveCache()
        
        logger.info("文档摄取管道已关闭")


from functools import lru_cache

@lru_cache(maxsize=1)
def _create_ingestion_pipeline(
    embeddingService=None,
    vectorStore=None,
    knowledgeService=None,
    ocrExtractor: Optional[Callable[[str], str]] = None,
) -> IngestionPipeline:
    return IngestionPipeline(
        embeddingService=embeddingService,
        vectorStore=vectorStore,
        knowledgeService=knowledgeService,
        ocrExtractor=ocrExtractor,
    )


def get_ingestion_pipeline(
    embeddingService=None,
    vectorStore=None,
    knowledgeService=None,
    ocrExtractor: Optional[Callable[[str], str]] = None,
    reset: bool = False
) -> IngestionPipeline:
    """
    获取摄取管道单例

    Args:
        embeddingService: 嵌入服务
        vectorStore: 向量存储
        knowledgeService: 知识服务
        reset: 是否重置

    Returns:
        IngestionPipeline实例
    """
    if reset:
        _create_ingestion_pipeline.cache_clear()
    pipeline = _create_ingestion_pipeline(embeddingService, vectorStore, knowledgeService, ocrExtractor)
    if knowledgeService is not None and pipeline.knowledgeService is None:
        pipeline.knowledgeService = knowledgeService
    if embeddingService is not None and pipeline.embeddingService is None:
        pipeline.embeddingService = embeddingService
    if vectorStore is not None and pipeline.vectorStore is None:
        pipeline.vectorStore = vectorStore
    if ocrExtractor is not None or getattr(pipeline, "ocrExtractor", None) is None:
        pipeline.set_ocr_extractor(ocrExtractor)
    return pipeline
